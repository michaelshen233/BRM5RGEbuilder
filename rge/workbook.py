"""Read cached Excel results. Never evaluate workbook formulas or instructions."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import re
import openpyxl
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from .commands import parse

COMMAND_START = re.compile(r"^(?:move|tween|spawn|create|explosion|teleport|bot|trigger|wait)(?:\s|\d|$)", re.I)


def read_workbook(payload):
    if len(payload) > 5 * 1024 * 1024:
        raise ValueError("Workbook must be smaller than 5 MB.")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 25 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise ValueError("Workbook expands beyond the 25 MB import limit.")
        formulas = openpyxl.load_workbook(BytesIO(payload), data_only=False, read_only=True, keep_links=False)
        cached = openpyxl.load_workbook(BytesIO(payload), data_only=True, read_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError, ParseError, DefusedXmlException, openpyxl.utils.exceptions.InvalidFileException) as error:
        raise ValueError("Could not read this .xlsx workbook.") from error
    entries, missing = [], []
    try:
        for sheet in cached:
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            if sheet.max_row > 20000 or sheet.max_column > 200:
                raise ValueError("Workbook dimensions exceed the import limit.")
            formula_rows = formulas[sheet.title].iter_rows()
            for row, originals in zip(sheet.iter_rows(), formula_rows):
                values = {i: c.value for i, c in enumerate(row, 1) if c.value is not None}
                for cell, original in zip(row, originals):
                    formula = original.value if original.data_type == "f" else None
                    if formula and cell.value is None:
                        missing.append(f"{sheet.title}!{original.coordinate}")
                    value = cell.value
                    if not isinstance(value, str) or not COMMAND_START.match(value.strip()):
                        continue
                    # Exclude headings like 'MOVE away' and 'TWEEN', not broken command records.
                    if (len(value.split()) < 2 and not formula) or value.strip() in ("MOVE away",):
                        continue
                    parsed = parse(value.strip())
                    label = values.get(1) if cell.column == 2 else (values.get(11) or values.get(16))
                    if not isinstance(label, str) or COMMAND_START.match(label) or len(label) > 100:
                        label = f'{parsed["kind"].replace("_", " ").title()} · {cell.coordinate}'
                    entries.append(dict(title=label, notes="Imported workbook formula result" if formula else "Imported workbook command", source=f"{sheet.title}!{cell.coordinate}", formula=formula, original=value, **parsed))
                    if len(entries) > 2000:
                        raise ValueError("Import supports up to 2,000 commands at a time.")
    finally:
        formulas.close()
        cached.close()
    return dict(entries=entries, missingFormulaResults=missing)
